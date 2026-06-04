from __future__ import annotations

import csv
from dataclasses import dataclass
import json
from pathlib import Path
from typing import Any


@dataclass(frozen=True)
class LoadedDocument:
    title: str
    text: str
    content_type: str
    metadata: dict[str, Any]


class DocumentLoadError(ValueError):
    pass


def load_document(path: str | Path) -> LoadedDocument:
    file_path = Path(path).expanduser()
    if not file_path.exists() or not file_path.is_file():
        raise DocumentLoadError(f"Document path does not exist: {file_path}")

    suffix = file_path.suffix.lower()
    if suffix in {".md", ".markdown"}:
        return _load_text(file_path, "text/markdown")
    if suffix == ".txt":
        return _load_text(file_path, "text/plain")
    if suffix == ".json":
        return _load_json(file_path)
    if suffix == ".csv":
        return _load_csv(file_path)
    if suffix == ".pdf":
        return _load_pdf(file_path)
    if suffix == ".docx":
        return _load_docx(file_path)

    raise DocumentLoadError(f"Unsupported document type: {suffix or 'unknown'}")


def _load_text(path: Path, content_type: str) -> LoadedDocument:
    text = path.read_text(encoding="utf-8")
    return LoadedDocument(path.stem, text, content_type, {"filename": path.name})


def _load_json(path: Path) -> LoadedDocument:
    data = json.loads(path.read_text(encoding="utf-8"))
    text = json.dumps(data, ensure_ascii=False, indent=2, sort_keys=True)
    return LoadedDocument(path.stem, text, "application/json", {"filename": path.name})


def _load_csv(path: Path) -> LoadedDocument:
    with path.open("r", encoding="utf-8", newline="") as handle:
        rows = [_normalize_csv_row(row) for row in csv.DictReader(handle)]
    if rows:
        header = ", ".join(rows[0].keys())
        body = "\n".join(json.dumps(row, ensure_ascii=False, sort_keys=True) for row in rows)
        text = f"Columns: {header}\n{body}"
    else:
        text = path.read_text(encoding="utf-8")
    return LoadedDocument(
        path.stem,
        text,
        "text/csv",
        {"filename": path.name, "row_count": len(rows)},
    )


def _normalize_csv_row(row: dict[Any, Any]) -> dict[str, Any]:
    normalized: dict[str, Any] = {}
    overflow = row.pop(None, None)
    for key, value in row.items():
        normalized[str(key)] = value
    if overflow:
        normalized["extra_values"] = overflow
    return normalized


def _load_pdf(path: Path) -> LoadedDocument:
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError as exc:
        raise DocumentLoadError("PDF support requires optional package 'pypdf'.") from exc

    reader = PdfReader(str(path))
    pages = [(page.extract_text() or "") for page in reader.pages]
    return LoadedDocument(
        path.stem,
        "\n\n".join(pages).strip(),
        "application/pdf",
        {"filename": path.name, "page_count": len(pages)},
    )


def _load_docx(path: Path) -> LoadedDocument:
    try:
        import docx  # type: ignore
    except ImportError as exc:
        raise DocumentLoadError("DOCX support requires optional package 'python-docx'.") from exc

    document = docx.Document(str(path))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs).strip()
    return LoadedDocument(
        path.stem,
        text,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        {"filename": path.name, "paragraph_count": len(document.paragraphs)},
    )
